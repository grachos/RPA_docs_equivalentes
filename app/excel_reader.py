"""Read and validate the purchase-order Excel template.

Template: one workbook, one sheet, header in row 1 with exactly these columns:
    A: ID POSEEDOR | B: Nombre del proveedor | C: orden de compra | D: Total
The workbook is copied to a temp file first (the user may have it open) and
opened read-only. All rows are read, regardless of active filters.
"""
from __future__ import annotations

import shutil
import tempfile
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from app.config import OC_LENGTH
from app.models import InvalidRow, Row
from app.text_utils import normalize_text, parse_total

TEMPLATE_HEADERS = ("ID POSEEDOR", "Nombre del proveedor", "orden de compra", "Total")


class TemplateError(ValueError):
    """The workbook does not follow the expected template."""


def _cell_to_digits(value: object) -> str:
    """Excel numbers come as int/float (11111111.0); text may carry spaces or dots."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    return str(value).strip().replace(" ", "").replace(".", "")


class _temp_copy:
    """Context manager that copies the workbook to a temp dir and cleans up."""

    def __init__(self, path: Path):
        self.path = Path(path)
        self._dir: str | None = None

    def __enter__(self) -> Path:
        if not self.path.exists():
            raise FileNotFoundError(f"No existe el archivo: {self.path}")
        self._dir = tempfile.mkdtemp(prefix="doc_soporte_")
        target = Path(self._dir) / self.path.name
        shutil.copy2(self.path, target)
        return target

    def __exit__(self, *exc) -> None:
        if self._dir:
            shutil.rmtree(self._dir, ignore_errors=True)


def check_headers(headers: tuple) -> None:
    """Raise TemplateError if row 1 does not match TEMPLATE_HEADERS (case/accents ignored)."""
    found = [normalize_text("" if h is None else str(h)) for h in headers]
    found += [""] * (len(TEMPLATE_HEADERS) - len(found))
    errors = []
    for col, (expected, got) in enumerate(zip(TEMPLATE_HEADERS, found)):
        if normalize_text(expected) != got:
            letter = "ABCD"[col]
            errors.append(f"columna {letter}: se esperaba {expected!r}, se encontró {got or '(vacío)'!r}")
    if errors:
        raise TemplateError("El Excel no sigue la plantilla: " + "; ".join(errors))


def validate_raw_row(excel_row: int, values: tuple) -> Row | InvalidRow:
    """Validate columns A-D of one row."""
    raw_id, raw_name, raw_oc, raw_total = (list(values) + [None] * 4)[:4]
    raw = {
        "id": "" if raw_id is None else str(raw_id),
        "name": "" if raw_name is None else str(raw_name),
        "oc": "" if raw_oc is None else str(raw_oc),
        "total": "" if raw_total is None else str(raw_total),
    }
    errors: list[str] = []

    provider_id = _cell_to_digits(raw_id)
    if not provider_id.isdigit():
        errors.append(f"ID no numérico: {raw['id']!r}")

    name = " ".join(raw["name"].split())
    if not name:
        errors.append("nombre vacío")

    oc = _cell_to_digits(raw_oc)
    if not (oc.isdigit() and len(oc) == OC_LENGTH):
        errors.append(f"OC debe tener {OC_LENGTH} dígitos: {raw['oc']!r}")

    total = 0
    try:
        total = parse_total(raw_total)
        if total <= 0:
            errors.append(f"total debe ser > 0: {raw['total']!r}")
    except ValueError as exc:
        errors.append(str(exc))

    if errors:
        return InvalidRow(excel_row=excel_row, errors=errors, raw=raw)
    return Row(excel_row=excel_row, provider_id=provider_id, provider_name=name, oc=oc, total=total)


def mark_duplicate_ocs(rows: list[Row]) -> tuple[list[Row], list[InvalidRow]]:
    """Every row whose OC appears more than once is rejected (all occurrences)."""
    by_oc: dict[str, list[Row]] = defaultdict(list)
    for row in rows:
        by_oc[row.oc].append(row)
    valid, invalid = [], []
    for row in rows:
        dupes = by_oc[row.oc]
        if len(dupes) > 1:
            lines = ", ".join(str(r.excel_row) for r in dupes)
            invalid.append(InvalidRow(
                excel_row=row.excel_row,
                errors=[f"OC {row.oc} duplicada (filas {lines})"],
                raw={"id": row.provider_id, "name": row.provider_name, "oc": row.oc,
                     "total": str(row.total)},
            ))
        else:
            valid.append(row)
    return valid, invalid


def read_orders(excel_path: Path) -> tuple[list[Row], list[InvalidRow], list[str]]:
    """Return (valid rows, invalid rows, warnings) from the workbook's single sheet."""
    warnings: list[str] = []
    with _temp_copy(excel_path) as copy:
        wb = load_workbook(copy, read_only=True, data_only=True)
        try:
            if len(wb.sheetnames) > 1:
                warnings.append(
                    f"El archivo tiene {len(wb.sheetnames)} hojas; solo se lee la primera "
                    f"({wb.sheetnames[0]!r})."
                )
            ws = wb[wb.sheetnames[0]]
            rows_iter = ws.iter_rows(max_col=len(TEMPLATE_HEADERS), values_only=True)
            header = next(rows_iter, None)
            if header is None:
                raise TemplateError("El Excel está vacío.")
            check_headers(header)
            parsed: list[Row] = []
            invalid: list[InvalidRow] = []
            for idx, values in enumerate(rows_iter, start=2):
                if all(v is None or str(v).strip() == "" for v in values):
                    continue
                result = validate_raw_row(idx, values)
                (parsed if isinstance(result, Row) else invalid).append(result)
        finally:
            wb.close()

    valid, dupes = mark_duplicate_ocs(parsed)
    return valid, sorted(invalid + dupes, key=lambda r: r.excel_row), warnings


def write_template(path: Path) -> Path:
    """Create an empty Excel with the expected template headers."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Ordenes"
    ws.append(list(TEMPLATE_HEADERS))
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell, width in zip(ws[1], (16, 42, 18, 16)):
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        ws.column_dimensions[cell.column_letter].width = width
    ws.freeze_panes = "A2"
    wb.save(path)
    return path
