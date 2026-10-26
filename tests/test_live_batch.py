"""REAL-mode row logic with fake MyE objects (no browser)."""
import io
import threading
from collections import defaultdict
import zipfile
from datetime import date
from pathlib import Path

import pytest

from app import live_batch as lb
from app.models import DocStatus, Row
from app.mye_documents import DocRow, DocumentsError, check_zip_number
from app.myeinvoice import DE_NUMBER_RE, MyeError, PreviewResult
from app.state import StateStore

ROW = Row(excel_row=2, provider_id="33333333", provider_name="ANA ROSA DIAZ MORA",
          oc="4600485726", total=1770000)


class FakeUi:
    def __init__(self, answer=True):
        self.answer, self.updates, self.asked = answer, [], 0

    def log(self, m):
        pass

    def row_update(self, oc, status, detail, tag):
        self.updates.append(status)

    def confirm(self, text, pdf, label=None):
        self.asked += 1
        assert "REGISTRAR" in text and ROW.oc in text
        return self.answer


class FakeMye:
    acquirer_name = acquirer_surnames = None

    def __init__(self, ok=True):
        self.ok, self.registered = ok, 0

    def register(self, row, observation, out_dir, approve):
        preview = PreviewResult(ok=self.ok, problems=[] if self.ok else ["total no aparece"])
        if not preview.ok or not approve(preview):
            return None, preview
        self.registered += 1
        return "DE252265", preview


class FakeDocs:
    def __init__(self, dupes=()):
        self.dupes = list(dupes)

    def possible_duplicates(self, *a):
        return self.dupes

    def wait_dian_ok(self, de, pid, since):
        return "Procesado Correctamente."


def make_batch(ui, mye, docs, auto=False):
    batch = object.__new__(lb.LiveBatch)
    batch.mye, batch.docs, batch.ui = mye, docs, ui
    batch.observation, batch.auto_confirm = "CORTE 103 22 octubre 2026", auto
    batch.cancel, batch.since = threading.Event(), date(2026, 8, 23)
    batch._mye_ready, batch.cadi = True, None
    batch.acquirers, batch._overlap, batch.created_oc = {}, False, {}
    batch.download = lambda row, de, out, obs=None: (Path(f"{de}.pdf"), Path(f"{de}.xml"))
    return batch


@pytest.fixture
def store(tmp_path):
    with StateStore(tmp_path / "s.db") as s:
        s.ensure_pending(ROW, "2026-10-22", "CORTE 103")
        yield s


def test_success_reaches_zip_downloaded(store, tmp_path):
    ui, mye = FakeUi(), FakeMye()
    assert make_batch(ui, mye, FakeDocs()).process_row(store, ROW, tmp_path) == "ok"
    st = store.get(ROW.oc, ROW.provider_id)
    assert st.status == DocStatus.ZIP_DOWNLOADED and st.de_number == "DE252265"
    assert st.dian_status == "Procesado Correctamente." and st.pdf_path.endswith("DE252265.pdf")
    assert ui.asked == 1 and mye.registered == 1


def test_user_declines_nothing_changes(store, tmp_path):
    ui, mye = FakeUi(answer=False), FakeMye()
    assert make_batch(ui, mye, FakeDocs()).process_row(store, ROW, tmp_path) == "omitida"
    assert store.get(ROW.oc, ROW.provider_id).status == DocStatus.PENDING
    assert mye.registered == 0


def test_bad_preview_never_registers_nor_asks(store, tmp_path):
    ui, mye = FakeUi(), FakeMye(ok=False)
    with pytest.raises(MyeError, match="no coincide con el Excel.*total no aparece"):
        make_batch(ui, mye, FakeDocs()).process_row(store, ROW, tmp_path)
    assert ui.asked == 0 and mye.registered == 0


def test_duplicate_stops_before_filling(store, tmp_path):
    mye = FakeMye()
    docs = FakeDocs([DocRow(0, "DE252264", "NIT 33333333", "Procesado", 1770000)])
    with pytest.raises(MyeError, match="duplicado.*DE252264"):
        make_batch(FakeUi(), mye, docs).process_row(store, ROW, tmp_path)
    assert mye.registered == 0


def test_resume_after_registered_does_not_register_again(store, tmp_path):
    store.advance(ROW.oc, ROW.provider_id, DocStatus.MYE_REGISTERED, de_number="DE252265")
    ui, mye = FakeUi(), FakeMye()
    assert make_batch(ui, mye, FakeDocs()).process_row(store, ROW, tmp_path) == "ok"
    assert mye.registered == 0 and ui.asked == 0
    assert store.get(ROW.oc, ROW.provider_id).status == DocStatus.ZIP_DOWNLOADED


def test_auto_confirm_skips_question(store, tmp_path):
    ui = FakeUi()
    make_batch(ui, FakeMye(), FakeDocs(), auto=True).process_row(store, ROW, tmp_path)
    assert ui.asked == 0


def test_de_number_regex():
    assert DE_NUMBER_RE.search('{"msg": "Se ha generado la Documento Soporte número DE252265"}').group(1) == "DE252265"
    assert DE_NUMBER_RE.search("idnumeracion=8759 DE25 Ultimo utilizado: 2264") is None


def test_zip_number_check(tmp_path):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("DE25_2265.zip", b"x")
    path = tmp_path / "a.zip"
    path.write_bytes(buf.getvalue())
    check_zip_number(path, "DE252265")
    with pytest.raises(DocumentsError):
        check_zip_number(path, "DE252266")


def test_three_rows_two_passes_and_error_isolated(tmp_path, monkeypatch):
    from app import config as cfg

    monkeypatch.setattr(cfg, "STATE_DB", tmp_path / "batch.db")
    rows = [ROW.model_copy(update={"oc": f"460000000{i}", "excel_row": i + 1, "total": 1000 * i})
            for i in (1, 2, 3)]
    with StateStore(cfg.STATE_DB) as s:
        for r in rows:
            s.ensure_pending(r, "2026-10-23", "CORTE 104")
    calls = []

    class Mye:
        acquirer_name = acquirer_surnames = None

        def register(self, row, observation, out_dir, approve):
            calls.append(("register", row.oc))
            if row.oc.endswith("2"):
                raise MyeError("proveedor no existe")
            return "DE25" + row.oc[-4:], PreviewResult(ok=True)

    class Docs:
        def possible_duplicates(self, *a):
            return []

        def wait_dian_ok(self, de, pid, since):
            calls.append(("dian", de))
            return "Procesado Correctamente."

    batch = make_batch(FakeUi(), Mye(), Docs(), auto=True)
    message = batch.run([(r, None) for r in rows], "CORTE 104")
    kinds = [c[0] for c in calls]
    assert kinds == ["register", "register", "register", "dian", "dian"]  # all registers first
    assert "2 terminadas" not in message and "2 con ZIP descargado, 1 con error" in message
    assert "Tiempo total" in message and "Registro MyE" in message
    with StateStore(cfg.STATE_DB) as s:
        assert [s.get(r.oc, r.provider_id).status for r in rows] == [
            DocStatus.ZIP_DOWNLOADED, DocStatus.FAILED, DocStatus.ZIP_DOWNLOADED]


def test_form_load_overlaps_duplicate_search(store, tmp_path):
    order = []

    class Mye(FakeMye):
        def start_open_form(self):
            order.append("start_form")

        def register(self, *a, **k):
            order.append("register")
            return super().register(*a, **k)

    class Docs(FakeDocs):
        def possible_duplicates(self, *a):
            order.append("duplicates")
            return []

    batch = make_batch(FakeUi(), Mye(), Docs(), auto=True)
    batch._overlap = True
    batch.process_row(store, ROW, tmp_path)
    assert order == ["start_form", "duplicates", "register"]


def test_times_are_stored_per_row_and_reported(tmp_path, monkeypatch):
    import json

    from openpyxl import load_workbook

    from app import config as cfg
    from app.report import write_report

    monkeypatch.setattr(cfg, "STATE_DB", tmp_path / "t.db")
    rows = [ROW.model_copy(update={"oc": f"460000001{i}", "total": 100 * i}) for i in (1, 2)]
    with StateStore(cfg.STATE_DB) as s:
        for r in rows:
            s.ensure_pending(r, "2026-10-23", "CORTE 105")
    batch = make_batch(FakeUi(), FakeMye(), FakeDocs(), auto=True)
    message = batch.run([(r, None) for r in rows], "CORTE 105")
    assert "OC 4600000011: Registro MyE" in message and "total" in message
    with StateStore(cfg.STATE_DB) as s:
        states = s.list(consolidado="CORTE 105")
        times = json.loads(states[0].timings)
        assert set(times) == {"Registro MyE", "Espera DIAN", "Descarga ZIP"}
        _, xlsx = write_report(states, tmp_path, "CORTE 105", batch.summary_rows())
    wb = load_workbook(xlsx)
    header = [c.value for c in wb["Reporte"][1]]
    assert header[-5:] == ["Registro MyE (s)", "Espera DIAN (s)", "Descarga ZIP (s)", "CADI (s)", "Total fila (s)"]
    assert wb["Reporte"].cell(2, len(header) - 4).value is not None
    resumen = {r[0].value: r[1].value for r in wb["Resumen"].iter_rows(min_row=2)}
    assert "Tiempo total del lote (min)" in resumen and resumen["Filas con tiempos medidos"] == 2


class Portal:
    """Fake CADI: invoices by number and by OC."""

    def __init__(self, invoices):
        self.invoices = invoices

    def find_by_oc(self, oc):
        return [i for i in self.invoices if i.oc == oc]

    def find_invoices(self, number):
        return [i for i in self.invoices if i.number == number]


def inv(number, oc, status="En revisión"):
    from app.cadi import CadiInvoice

    return CadiInvoice("2340" + number[-2:], number, "FAC ELEC - TERCERIZACION", ROW.provider_id, "X",
                       oc, status, "1")


def batch_with_cadi(invoices, candidates):
    from types import SimpleNamespace

    batch = make_batch(FakeUi(), FakeMye(), FakeDocs(candidates), auto=True)
    batch.cadi = SimpleNamespace(cadi=Portal(invoices))
    batch.times, batch.row_times, batch._pending_times = defaultdict(list), defaultdict(dict), []
    return batch


SAME_PROVIDER_AMOUNT = [DocRow(0, "DE251961", "NIT 33333333-1", "Procesado", 1770000)]


def test_same_provider_and_amount_but_other_oc_is_processed(store, tmp_path):
    # User rule: DE251961 is an older service (another OC) -> a new DE must be created.
    batch = batch_with_cadi([inv("DE251961", "4600400000", "Egreso")], SAME_PROVIDER_AMOUNT)
    assert batch.register_stage(store, ROW, tmp_path) == "continue"
    assert store.get(ROW.oc, ROW.provider_id).de_number == "DE252265"


def test_oc_already_in_cadi_is_duplicate(store, tmp_path):
    batch = batch_with_cadi([inv("DE252264", ROW.oc)], [])
    with pytest.raises(MyeError, match=f"OC {ROW.oc} ya está registrada en CADI: factura 234064 .DE252264"):
        batch.process_row(store, ROW, tmp_path)
    assert store.get(ROW.oc, ROW.provider_id).status == DocStatus.PENDING


def test_same_amount_document_with_unknown_oc_stops_for_review(store, tmp_path):
    batch = batch_with_cadi([], SAME_PROVIDER_AMOUNT)  # DE251961 is not in CADI
    with pytest.raises(MyeError, match="Revisar: DE251961.*no está en CADI o no tiene OC"):
        batch.process_row(store, ROW, tmp_path)


def test_document_created_in_this_batch_uses_its_known_oc(store, tmp_path):
    batch = batch_with_cadi([], [DocRow(-1, "DE252290", "NIT 33333333", "", 1770000)])
    batch.created_oc["DE252290"] = "4600499999"  # other OC registered earlier in this run
    assert batch.register_stage(store, ROW, tmp_path) == "continue"
    assert store.get(ROW.oc, ROW.provider_id).de_number == "DE252265"
