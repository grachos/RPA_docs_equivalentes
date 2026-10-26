import pytest
from openpyxl import Workbook, load_workbook

from app.excel_reader import (
    TEMPLATE_HEADERS,
    TemplateError,
    check_headers,
    read_orders,
    validate_raw_row,
    write_template,
)
from app.models import InvalidRow, Row


def _save(path, rows, headers=TEMPLATE_HEADERS, extra_sheet=False):
    wb = Workbook()
    ws = wb.active
    ws.append(list(headers))
    for r in rows:
        ws.append(r)
    if extra_sheet:
        wb.create_sheet("OTRA")
    wb.save(path)
    return path


@pytest.fixture
def excel(tmp_path):
    path = _save(tmp_path / "ordenes.xlsx", [
        [11111111, "ANA MARIA PEREZ LOPEZ", 4600485727, 3250000],
        ["1.023.456", "JUAN  PEREZ", "4600485558", "1.200.000"],
        [None, None, None, None],                      # empty: skipped
        ["ABC", "MALO", 123, 0],                       # invalid
        [111, "DUP UNO", 4600000001, 10],              # duplicate OC
        [222, "DUP DOS", 4600000001, 20],              # duplicate OC
    ])
    wb = load_workbook(path)
    wb.active.auto_filter.ref = "A1:D7"
    wb.active.row_dimensions[3].hidden = True  # hidden by filter: must still be read
    wb.save(path)
    return path


def test_read_orders_valid_and_invalid(excel):
    rows, invalid, warnings = read_orders(excel)
    assert [r.oc for r in rows] == ["4600485727", "4600485558"]
    assert rows[0].total == 3250000 and rows[0].provider_id == "11111111"
    assert rows[1].provider_id == "1023456" and rows[1].provider_name == "JUAN PEREZ"
    assert [r.excel_row for r in invalid] == [5, 6, 7]
    assert any("duplicada" in e for e in invalid[1].errors)
    assert warnings == []


def test_headers_case_and_accents_ignored(tmp_path):
    path = _save(tmp_path / "x.xlsx", [[1, "A", 4600000001, 5]],
                 headers=("id poseedor", "NOMBRE DEL PROVEEDOR", "Orden de Compra", "TOTAL"))
    rows, _, _ = read_orders(path)
    assert len(rows) == 1


def test_missing_provider_name_header(tmp_path):
    path = _save(tmp_path / "x.xlsx", [], headers=("ID POSEEDOR", None, "orden de compra", "Total"))
    with pytest.raises(TemplateError, match="columna B"):
        read_orders(path)


def test_check_headers_short_row():
    with pytest.raises(TemplateError):
        check_headers(("ID POSEEDOR", "Nombre del proveedor"))


def test_extra_sheets_warn_and_read_first(tmp_path):
    path = _save(tmp_path / "x.xlsx", [[1, "A", 4600000001, 5]], extra_sheet=True)
    rows, _, warnings = read_orders(path)
    assert len(rows) == 1 and "solo se lee la primera" in warnings[0]


def test_template_roundtrip(tmp_path):
    path = write_template(tmp_path / "plantilla.xlsx")
    rows, invalid, _ = read_orders(path)
    assert rows == [] and invalid == []


def test_validate_raw_row_errors():
    result = validate_raw_row(9, ("819004116-5", "", "46004855", "-5"))
    assert isinstance(result, InvalidRow)
    assert len(result.errors) == 4


def test_validate_raw_row_ok():
    result = validate_raw_row(2, (11111111.0, "X", 4600485727.0, 100.0))
    assert isinstance(result, Row) and result.oc == "4600485727"
